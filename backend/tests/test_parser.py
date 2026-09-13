from app.services.extraction.models import (
    ExtractedDocument,
    ExtractedElement,
    Provenance,
)
from app.services.parser.dates import (
    extract_date_from_text,
    parse_date_string,
)
from app.services.parser.parser import SemanticResumeParser
from app.services.parser.sections import (
    SectionType,
    classify_heading,
)

# =========================================================================
# 1. ExtractedDocument and Provenance Tests
# =========================================================================


def test_extracted_document_and_numeric_provenance():
    text = "Aayush Jadhav\nSoftware Engineer\n\nEDUCATION\nIIT Bombay\n\nEXPERIENCE\nGoogle"
    doc = ExtractedDocument.from_text(text, page_number=1)

    assert len(doc.elements) == 3
    for idx, elem in enumerate(doc.elements, start=1):
        assert elem.element_id == idx
        assert elem.page_number == 1
        assert elem.order == idx
        prov = elem.provenance
        assert isinstance(prov, Provenance)
        assert prov.page == 1
        assert prov.element == idx
        assert prov.order == idx


def test_extracted_element_numeric_only_invariants():
    elem = ExtractedElement(
        element_id=42,
        page_number=2,
        order=42,
        element_type="paragraph",
        text="Test text",
    )
    prov = elem.provenance
    assert isinstance(prov.page, int)
    assert isinstance(prov.element, int)
    assert isinstance(prov.order, int)
    assert prov.page == 2
    assert prov.element == 42
    assert prov.order == 42


# =========================================================================
# 2. Date Handling, Normalization, Ambiguity, and Precision Tests
# =========================================================================


def test_date_normalization_standard_month_year():
    # Full month names
    d1 = parse_date_string("January 2024")
    assert d1 is not None
    assert d1.original == "January 2024"
    assert d1.normalized == "01/2024"
    assert d1.start_date == "01/2024"
    assert not d1.is_present

    # Abbreviated month names
    d2 = parse_date_string("Aug 2023")
    assert d2 is not None
    assert d2.original == "Aug 2023"
    assert d2.normalized == "08/2023"

    # Numeric MM/YYYY
    d3 = parse_date_string("05/2022")
    assert d3 is not None
    assert d3.original == "05/2022"
    assert d3.normalized == "05/2022"


def test_date_normalization_ranges_and_present():
    # Month-Year range
    d_range = parse_date_string("January 2024 - March 2025")
    assert d_range is not None
    assert d_range.original == "January 2024 - March 2025"
    assert d_range.normalized == "01/2024 – 03/2025"
    assert d_range.start_date == "01/2024"
    assert d_range.end_date == "03/2025"
    assert not d_range.is_present

    # Ongoing / Present
    d_pres = parse_date_string("July 2023 – Present")
    assert d_pres is not None
    assert d_pres.original == "July 2023 – Present"
    assert d_pres.normalized == "07/2023 – Present"
    assert d_pres.start_date == "07/2023"
    assert d_pres.end_date == "Present"
    assert d_pres.is_present is True

    # Current / Ongoing variations
    d_curr = parse_date_string("01/2022 to Current")
    assert d_curr is not None
    assert d_curr.normalized == "01/2022 – Present"
    assert d_curr.is_present is True


def test_date_partial_dates_preserve_granularity():
    # Year only: 2024 must NOT become 01/2024
    d_yr = parse_date_string("2024")
    assert d_yr is not None
    assert d_yr.original == "2024"
    assert d_yr.normalized == "2024"
    assert d_yr.start_date == "2024"

    # Year range: 2020 - 2024
    d_yr_range = parse_date_string("2020 - 2024")
    assert d_yr_range is not None
    assert d_yr_range.original == "2020 - 2024"
    assert d_yr_range.normalized == "2020 – 2024"
    assert d_yr_range.start_date == "2020"
    assert d_yr_range.end_date == "2024"


def test_date_ambiguous_dates_never_guessed():
    # 03/04/2024 (Ambiguous whether March 4 or April 3)
    d_ambig = parse_date_string("03/04/2024")
    assert d_ambig is not None
    assert d_ambig.original == "03/04/2024"
    assert d_ambig.normalized is None
    assert d_ambig.start_date is None
    assert d_ambig.end_date is None

    # Range of ambiguous dates
    d_ambig_range = parse_date_string("03/04/2022 - 05/06/2023")
    assert d_ambig_range is not None
    assert d_ambig_range.original == "03/04/2022 - 05/06/2023"
    assert d_ambig_range.normalized is None


def test_date_unambiguous_day_month_year():
    # 25/04/2024 (25 > 12 -> must be Day=25, Month=04)
    d_unambig = parse_date_string("25/04/2024")
    assert d_unambig is not None
    assert d_unambig.original == "25/04/2024"
    assert d_unambig.normalized == "04/2024"


def test_extract_date_from_text_line():
    line = "Software Engineer | Acme Corp (Jan 2021 - Present)"
    d, clean = extract_date_from_text(line)
    assert d is not None
    assert d.normalized == "01/2021 – Present"
    assert d.is_present is True
    assert "Acme Corp" in clean
    assert "Jan 2021" not in clean


# =========================================================================
# 3. Section Classification and Synonym Tests
# =========================================================================


def test_section_classification_synonyms():
    # Experience synonyms
    assert classify_heading("WORK EXPERIENCE")[0] == SectionType.EXPERIENCE
    assert classify_heading("Professional Experience")[0] == SectionType.EXPERIENCE
    assert classify_heading("Employment History")[0] == SectionType.EXPERIENCE
    assert classify_heading("Career History")[0] == SectionType.EXPERIENCE
    assert classify_heading("Internships")[0] == SectionType.EXPERIENCE

    # Education synonyms
    assert classify_heading("EDUCATION")[0] == SectionType.EDUCATION
    assert classify_heading("Academic Background")[0] == SectionType.EDUCATION
    assert classify_heading("Educational Qualifications")[0] == SectionType.EDUCATION
    assert classify_heading("Academics")[0] == SectionType.EDUCATION

    # Projects synonyms
    assert classify_heading("PROJECTS")[0] == SectionType.PROJECTS
    assert classify_heading("Personal Projects")[0] == SectionType.PROJECTS
    assert classify_heading("Key Technical Projects")[0] == SectionType.PROJECTS

    # Skills synonyms
    assert classify_heading("TECHNICAL SKILLS")[0] == SectionType.SKILLS
    assert classify_heading("Core Competencies")[0] == SectionType.SKILLS
    assert classify_heading("Tech Stack")[0] == SectionType.SKILLS

    # Certifications synonyms
    assert (
        classify_heading("Licenses & Certifications")[0] == SectionType.CERTIFICATIONS
    )
    assert classify_heading("Courses & Certifications")[0] == SectionType.CERTIFICATIONS

    # Awards synonyms
    assert classify_heading("Honors & Awards")[0] == SectionType.AWARDS
    assert classify_heading("Achievements")[0] == SectionType.AWARDS

    # Publications synonyms
    assert classify_heading("Research Publications")[0] == SectionType.PUBLICATIONS
    assert classify_heading("Patents & Publications")[0] == SectionType.PUBLICATIONS

    # Volunteering & Organizations
    assert classify_heading("Volunteer Experience")[0] == SectionType.VOLUNTEERING
    assert (
        classify_heading("Positions of Responsibility")[0] == SectionType.ORGANIZATIONS
    )


def test_section_classification_preserves_original_title():
    sec_type, orig_title = classify_heading("1. PROFESSIONAL WORK EXPERIENCE:")
    assert sec_type == SectionType.EXPERIENCE
    assert orig_title == "1. PROFESSIONAL WORK EXPERIENCE:"


def test_section_classification_unknown_section_falls_back_to_additional():
    sec_type, orig_title = classify_heading("PATENTS & INVENTIONS 2024")
    # If not in exact dictionary, falls back to additional
    sec_type2, orig_title2 = classify_heading("CUSTOM PROPRIETARY CLEARANCES")
    assert sec_type2 == SectionType.ADDITIONAL
    assert orig_title2 == "CUSTOM PROPRIETARY CLEARANCES"


# =========================================================================
# 4. Scenario 1: Indian Engineering Student Resume
# =========================================================================


def test_scenario_indian_engineering_student_resume():
    raw_text = """
Rahul Sharma
rahul.sharma@iitb.ac.in | +91 9876543210 | Mumbai, Maharashtra, India
https://linkedin.com/in/rahulsharma | https://github.com/rahul-sharma

CAREER OBJECTIVE
Enthusiastic computer science undergraduate seeking a software engineering role with expertise in distributed systems and backend architecture.

EDUCATION
Indian Institute of Technology Bombay | Mumbai, India
B.Tech in Computer Science and Engineering (2021 - 2025)
CGPA: 9.42 / 10.0

Delhi Public School, R.K. Puram | New Delhi, India
Class XII (CBSE) - 2021
Percentage: 97.4%

EXPERIENCE
Software Engineering Intern at Zepto, Bengaluru
May 2024 - July 2024
• Architected real-time order tracking microservice handling 50k RPM with sub-50ms latency.
• Implemented Redis distributed caching layer reducing database load by 40%.

TECHNICAL SKILLS
Programming Languages: Python, C++, Go, TypeScript, Java
Frameworks & Libraries: FastAPI, Django, React, Next.js, Node.js
Databases & Cloud: PostgreSQL, Redis, MongoDB, Docker, AWS (S3, EC2)

PROJECTS
NextRound - AI Placement Readiness Platform (Jan 2024 - Mar 2024)
Technologies: Python, FastAPI, PostgreSQL, Docker
https://github.com/rahul-sharma/nextround
• Built automated resume parser with deterministic provenance tracking.
• Designed interactive skill gap analysis engine for placement students.

POSITIONS OF RESPONSIBILITY
Technical Secretary | Web and Coding Club, IIT Bombay (Aug 2023 - Present)
• Led a team of 40 organizers conducting annual hackathons with 2000+ participants.

AWARDS & ACHIEVEMENTS
Winner - Smart India Hackathon 2023 (Sep 2023)
• Awarded 1st prize out of 500+ competing engineering teams nationwide.
"""

    parser = SemanticResumeParser()
    resume = parser.parse_text(raw_text)

    # 1. Identity
    assert resume.identity.name == "Rahul Sharma"
    assert resume.identity.email == "rahul.sharma@iitb.ac.in"
    assert resume.identity.phone == "+91 9876543210"
    assert "Mumbai" in (resume.identity.location or "")
    assert len(resume.identity.links) == 2

    # 2. Summary
    assert resume.summary is not None
    assert resume.summary.original_title == "CAREER OBJECTIVE"
    assert "Enthusiastic computer science undergraduate" in resume.summary.text

    # 3. Education
    assert len(resume.education) == 2
    iit = resume.education[0]
    assert iit.institution == "Indian Institute of Technology Bombay"
    assert iit.degree == "B.Tech"
    assert iit.field_of_study == "Computer Science and Engineering"
    assert iit.dates is not None
    assert iit.dates.normalized == "2021 – 2025"
    assert any("9.42" in desc for desc in iit.description)

    # 4. Experience
    assert len(resume.experience) == 1
    exp = resume.experience[0]
    assert exp.organization == "Zepto"
    assert exp.role == "Software Engineering Intern"
    assert exp.location == "Bengaluru"
    assert exp.dates is not None
    assert exp.dates.normalized == "05/2024 – 07/2024"
    assert len(exp.description) == 2
    assert "50k RPM" in exp.description[0]

    # 5. Skills
    assert len(resume.skills) >= 3
    lang_cat = next(c for c in resume.skills if c.category == "Programming Languages")
    assert "Python" in lang_cat.skills
    assert "Go" in lang_cat.skills

    # 6. Projects
    assert len(resume.projects) == 1
    proj = resume.projects[0]
    assert "NextRound" in (proj.name or "")
    assert proj.dates is not None
    assert proj.dates.normalized == "01/2024 – 03/2024"
    assert "FastAPI" in proj.technologies
    assert len(proj.links) == 1
    assert "github.com" in proj.links[0].url

    # 7. Organizations
    assert len(resume.organizations) == 1
    org = resume.organizations[0]
    assert "Web and Coding Club" in (org.role or org.organization or "")
    assert org.dates is not None
    assert org.dates.is_present is True

    # 8. Awards
    assert len(resume.awards) == 1
    award = resume.awards[0]
    assert "Smart India Hackathon 2023" in award.title
    assert award.dates is not None
    assert award.dates.normalized == "09/2023"


# =========================================================================
# 5. Scenario 2: US-Style Professional Resume
# =========================================================================


def test_scenario_us_professional_resume():
    raw_text = """
Sarah Jenkins
sarah.jenkins@example.com | (415) 555-0199 | San Francisco, CA
https://linkedin.com/in/sarahjenkins | https://sarahjenkins.dev

PROFESSIONAL SUMMARY
Senior Backend Engineer with 8+ years of experience designing fault-tolerant cloud infrastructures.

WORK EXPERIENCE
Senior Software Engineer at Stripe, San Francisco, CA
March 2021 - Present
• Scaled payment processing pipeline to support $10B+ in annual transaction volume.
• Mentored 6 junior and mid-level software engineers on high-concurrency Go services.

Staff Software Engineer at Twilio, San Francisco, CA
June 2018 - February 2021
• Reduced API gateway p99 latency by 35% through connection pooling optimizations.

EDUCATION
Stanford University | Stanford, CA
Bachelor of Science in Computer Science (2014 - 2018)

CERTIFICATIONS
AWS Certified Solutions Architect - Professional | Amazon Web Services (Jan 2023)
Credential ID: AWS-9928172
"""

    parser = SemanticResumeParser()
    resume = parser.parse_text(raw_text)

    # Identity
    assert resume.identity.name == "Sarah Jenkins"
    assert resume.identity.email == "sarah.jenkins@example.com"
    assert resume.identity.phone == "(415) 555-0199"
    assert resume.identity.location == "San Francisco, CA"

    # Experience
    assert len(resume.experience) == 2
    stripe = resume.experience[0]
    assert stripe.organization == "Stripe"
    assert stripe.role == "Senior Software Engineer"
    assert stripe.location == "San Francisco, CA"
    assert stripe.dates is not None
    assert stripe.dates.normalized == "03/2021 – Present"
    assert stripe.dates.is_present is True

    twilio = resume.experience[1]
    assert twilio.organization == "Twilio"
    assert twilio.dates is not None
    assert twilio.dates.normalized == "06/2018 – 02/2021"

    # Education
    assert len(resume.education) == 1
    assert resume.education[0].institution == "Stanford University"
    assert resume.education[0].degree == "Bachelor of Science"

    # Certifications
    assert len(resume.certifications) == 1
    cert = resume.certifications[0]
    assert "AWS Certified Solutions Architect" in cert.name
    assert cert.issuer == "Amazon Web Services"
    assert cert.credential_id == "AWS-9928172"
    assert cert.dates is not None
    assert cert.dates.normalized == "01/2023"


# =========================================================================
# 6. Scenario 3: Academic / Research CV
# =========================================================================


def test_scenario_academic_research_cv():
    raw_text = """
Dr. Elena Rostova
elena.rostova@cam.ac.uk | Cambridge, UK

ACADEMIC BACKGROUND
University of Cambridge | Cambridge, UK
Doctor of Philosophy in Computer Science (2019 - 2023)
Thesis: Robust Representation Learning in Sparse Graph Neural Networks

PUBLICATIONS
Adaptive Graph Attention Networks for Molecular Discovery (2023)
https://doi.org/10.1145/3534678
• Published in IEEE Transactions on Pattern Analysis and Machine Intelligence.

FELLOWSHIPS & AWARDS
Turing Doctoral Fellowship | Alan Turing Institute (2019)
• Fully funded 4-year doctoral research fellowship.

LANGUAGES
English (Native)
Russian (Native)
French (Fluent)
"""

    parser = SemanticResumeParser()
    resume = parser.parse_text(raw_text)

    # Identity
    assert resume.identity.name == "Dr. Elena Rostova"
    assert resume.identity.email == "elena.rostova@cam.ac.uk"

    # Education
    assert len(resume.education) == 1
    assert resume.education[0].degree == "Doctor of Philosophy"
    assert resume.education[0].institution == "University of Cambridge"

    # Publications
    assert len(resume.publications) == 1
    pub = resume.publications[0]
    assert "Adaptive Graph Attention Networks" in pub.title
    assert pub.dates is not None
    assert pub.dates.normalized == "2023"
    assert len(pub.links) == 1
    assert "doi.org" in pub.links[0].url

    # Awards
    assert len(resume.awards) == 1
    assert "Turing Doctoral Fellowship" in resume.awards[0].title
    assert resume.awards[0].dates is not None
    assert resume.awards[0].dates.normalized == "2019"

    # Languages
    assert len(resume.languages) == 3
    eng = next(lang for lang in resume.languages if lang.language == "English")
    assert eng.proficiency == "Native"
    fr = next(lang for lang in resume.languages if lang.language == "French")
    assert fr.proficiency == "Fluent"


# =========================================================================
# 7. Scenario 4: Custom / Unknown Sections Preserved
# =========================================================================


def test_scenario_custom_unknown_sections_preserved():
    raw_text = """
John Developer
john@dev.com

WORK EXPERIENCE
Software Developer at TechCorp (2022 - 2024)

SECURITY CLEARANCES & SPECIAL ACCESS
• Top Secret / SCI Cleared (Active)
• Special Access Program (SAP) Nominated 2023

MILITARY SERVICE
Sergeant | US Army Signal Corps (2016 - 2020)
• Managed secure tactical communication networks.
"""

    parser = SemanticResumeParser()
    resume = parser.parse_text(raw_text)

    assert len(resume.experience) == 1
    # Unknown sections should be preserved under additional_sections
    assert len(resume.additional_sections) == 2
    titles = [s.original_title for s in resume.additional_sections]
    assert "SECURITY CLEARANCES & SPECIAL ACCESS" in titles
    assert "MILITARY SERVICE" in titles
    sec1 = next(
        s
        for s in resume.additional_sections
        if s.original_title == "SECURITY CLEARANCES & SPECIAL ACCESS"
    )
    assert any("Top Secret" in line for line in sec1.content)


# =========================================================================
# 8. Scenario 5: Minimal Resume
# =========================================================================


def test_scenario_minimal_resume():
    raw_text = """
Jane Minimalist
jane.min@example.org
"""
    parser = SemanticResumeParser()
    resume = parser.parse_text(raw_text)

    assert resume.identity.name == "Jane Minimalist"
    assert resume.identity.email == "jane.min@example.org"
    assert resume.identity.phone is None
    assert resume.identity.location is None
    assert resume.experience == []
    assert resume.education == []
    assert resume.skills == []
    assert resume.additional_sections == []


# =========================================================================
# 9. Scenario 6: Resume with Ambiguous Dates
# =========================================================================


def test_scenario_ambiguous_dates_preserved_without_guessing():
    raw_text = """
Alex Mercer
alex@example.com

WORK EXPERIENCE
Software Engineer at Delta Labs (04/05/2021 - 06/07/2023)
• Developed features in Python.
"""
    parser = SemanticResumeParser()
    resume = parser.parse_text(raw_text)

    assert len(resume.experience) == 1
    exp = resume.experience[0]
    assert exp.dates is not None
    # 04/05/2021 is ambiguous (April 5 or May 4) -> normalized must be None, original preserved!
    assert exp.dates.normalized is None
    assert exp.dates.original == "04/05/2021 - 06/07/2023"


# =========================================================================
# 10. Source Order Preservation Tests
# =========================================================================


def test_source_order_preservation_not_sorted():
    raw_text = """
Candidate Name
cand@example.com

WORK EXPERIENCE
Junior Engineer at Alpha (2018 - 2019)
Senior Engineer at Beta (2022 - Present)
Intern at Gamma (2016 - 2017)

TECHNICAL SKILLS
ZooKeeper, Ansible, Python, C++, React
"""
    parser = SemanticResumeParser()
    resume = parser.parse_text(raw_text)

    # Document order was Alpha, Beta, Gamma (not chronological Beta, Alpha, Gamma)
    assert len(resume.experience) == 3
    assert resume.experience[0].organization == "Alpha"
    assert resume.experience[1].organization == "Beta"
    assert resume.experience[2].organization == "Gamma"

    # Skills order must preserve document order (not alphabetical)
    assert len(resume.skills) >= 1
    skill_list = resume.skills[0].skills
    assert skill_list == ["ZooKeeper", "Ansible", "Python", "C++", "React"]


# =========================================================================
# 11. Extractor Pipeline Integration Tests (PDF, DOCX, Service)
# =========================================================================


def test_pdf_extractor_to_parser_pipeline(tmp_path):
    from app.services.extraction import PDFExtractor
    from tests.test_extraction import create_pdf_bytes

    pdf_bytes = create_pdf_bytes(
        [
            "Alice Wonder\nalice@example.com | (555) 000-1111\n\nEDUCATION\nMIT\nB.S. in Computer Science (2020 - 2024)",
            "WORK EXPERIENCE\nSoftware Engineer at Startup Co\nJan 2024 - Present\n* Built high throughput event queues.",
        ]
    )

    pdf_path = tmp_path / "alice_resume.pdf"
    pdf_path.write_bytes(pdf_bytes)

    extractor = PDFExtractor()
    extracted_doc = extractor.extract_document(pdf_path)

    assert len(extracted_doc.elements) >= 2
    # Verify page provenance is present on extracted elements
    pages = [e.page_number for e in extracted_doc.elements]
    assert 1 in pages
    assert 2 in pages

    parser = SemanticResumeParser()
    canonical = parser.parse(extracted_doc)

    assert canonical.identity.name == "Alice Wonder"
    assert canonical.identity.email == "alice@example.com"
    assert len(canonical.education) == 1
    assert canonical.education[0].institution == "MIT"
    assert canonical.education[0].degree == "B.S."
    assert len(canonical.experience) == 1
    assert canonical.experience[0].organization == "Startup Co"
    assert canonical.experience[0].dates is not None
    assert canonical.experience[0].dates.is_present is True


def test_docx_extractor_to_parser_pipeline(tmp_path):
    import docx

    from app.services.extraction import DOCXExtractor

    docx_path = tmp_path / "bob_resume.docx"
    doc = docx.Document()
    doc.add_paragraph("Bob Martin")
    doc.add_paragraph("bob@cleanarchitecture.com | +1 800 555 0100")
    doc.add_paragraph("CAREER SUMMARY")
    doc.add_paragraph("Author and software craftsmanship pioneer.")
    doc.add_paragraph("WORK EXPERIENCE")
    doc.add_paragraph("Principal Consultant at Clean Coders")
    doc.add_paragraph("2010 - Present")
    doc.add_paragraph("• Created software architecture educational videos.")
    doc.save(docx_path)

    extractor = DOCXExtractor()
    extracted_doc = extractor.extract_document(docx_path)

    parser = SemanticResumeParser()
    canonical = parser.parse(extracted_doc)

    assert canonical.identity.name == "Bob Martin"
    assert canonical.identity.email == "bob@cleanarchitecture.com"
    assert canonical.summary is not None
    assert "software craftsmanship pioneer" in canonical.summary.text
    assert len(canonical.experience) == 1
    assert canonical.experience[0].organization == "Clean Coders"
    assert canonical.experience[0].dates is not None
    assert canonical.experience[0].dates.is_present is True


def test_no_hallucination_guarantees():
    """Verify parser never invents unprovided data."""
    raw_text = """
Unknown Person
user@test.org

PROJECTS
Mini Project
• Built a simple CLI tool.
"""
    parser = SemanticResumeParser()
    canonical = parser.parse_text(raw_text)

    # Missing fields must remain None / empty, never fabricated
    assert canonical.identity.phone is None
    assert canonical.identity.location is None
    assert canonical.identity.links == []
    assert canonical.summary is None
    assert canonical.education == []
    assert canonical.experience == []
    assert canonical.skills == []
    assert canonical.certifications == []
    assert canonical.awards == []
    assert canonical.publications == []
    assert canonical.languages == []
    assert canonical.volunteering == []
    assert canonical.organizations == []
    assert len(canonical.projects) == 1
    assert canonical.projects[0].name == "Mini Project"
    assert canonical.projects[0].dates is None
    assert canonical.projects[0].technologies == []

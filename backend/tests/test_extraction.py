import uuid
from pathlib import Path
from unittest.mock import MagicMock, patch

import docx
import pytest

from app.models.resume import Resume, ResumeStatus
from app.services.extraction import (
    DocumentContentError,
    DocumentExtractionService,
    DocumentReadError,
    DOCXExtractor,
    ExtractionIssue,
    ExtractionResult,
    ExtractionStatus,
    ExtractorRegistry,
    PDFExtractor,
    UnsupportedDocumentError,
    normalize_text,
)
from app.services.storage import LocalStorageService, StoragePathTraversalError


def create_pdf_bytes(pages_text: list[str]) -> bytes:
    """Constructs a valid PDF binary byte stream with the specified page texts."""
    num_pages = len(pages_text)
    kids_refs = " ".join(f"{3 + i} 0 R" for i in range(num_pages))
    font_obj_idx = 3 + num_pages
    content_start_idx = font_obj_idx + 1

    catalog = "1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj"
    pages = (
        f"2 0 obj\n<< /Type /Pages /Kids [{kids_refs}] /Count {num_pages} >>\nendobj"
    )
    font = f"{font_obj_idx} 0 obj\n<< /Type /Font /Subtype /Type1 /Name /F1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>\nendobj"

    page_objs: list[str] = []
    content_objs: list[bytes] = []
    for i, txt in enumerate(pages_text):
        c_idx = content_start_idx + i
        p_idx = 3 + i
        if txt:
            lines = txt.split("\n")
            stream_parts = ["BT", "/F1 12 Tf", "72 700 Td"]
            for line_idx, line in enumerate(lines):
                safe_line = line.replace("(", "\\(").replace(")", "\\)")
                if line_idx > 0:
                    stream_parts.append("0 -15 Td")
                stream_parts.append(f"({safe_line}) Tj")
            stream_parts.append("ET\n")
            stream_content = "\n".join(stream_parts).encode("latin1")
        else:
            stream_content = b""

        p_obj = f"{p_idx} 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_obj_idx} 0 R >> >> /Contents {c_idx} 0 R >>\nendobj"
        c_obj = (
            f"{c_idx} 0 obj\n<< /Length {len(stream_content)} >>\nstream\n".encode(
                "latin1"
            )
            + stream_content
            + b"endstream\nendobj"
        )
        page_objs.append(p_obj)
        content_objs.append(c_obj)

    all_objs = (
        [catalog.encode("latin1"), pages.encode("latin1")]
        + [p.encode("latin1") for p in page_objs]
        + [font.encode("latin1")]
        + content_objs
    )

    body = b"%PDF-1.4\n"
    offsets = [0]
    for obj in all_objs:
        offsets.append(len(body))
        body += obj + b"\n"

    startxref = len(body)
    xref = f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode("latin1")
    for off in offsets[1:]:
        xref += f"{off:010d} 00000 n \n".encode("latin1")
    trailer = f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{startxref}\n%%EOF\n".encode(
        "latin1"
    )
    return body + xref + trailer


# ==========================================
# 1. TEXT NORMALIZATION TESTS
# ==========================================


def test_normalize_text_crlf_and_cr() -> None:
    raw = "Line 1\r\nLine 2\rLine 3\nLine 4"
    normalized = normalize_text(raw)
    assert normalized == "Line 1\nLine 2\nLine 3\nLine 4"


def test_normalize_text_trailing_and_excessive_whitespace() -> None:
    raw = "   Header with spaces   \t \n\n\n\n   Paragraph text with \t multiple    spaces.   \n\n\n"
    normalized = normalize_text(raw)
    assert normalized == "Header with spaces\n\nParagraph text with multiple spaces."


def test_normalize_text_unicode_and_special_characters() -> None:
    # Non-breaking space \u00a0, soft hyphen \xad, zero-width space \u200b
    raw = "Software\u00a0Engineer\nNode.\u200bjs\nResume\xadTitle"
    normalized = normalize_text(raw)
    assert normalized == "Software Engineer\nNode.js\nResumeTitle"


def test_normalize_text_preserves_technical_tokens() -> None:
    raw = "Proficient in C++, C#, .NET 8, Node.js, JavaScript, and SQL.\nWebsite: https://nextround.in\nEmail: contact@nextround.in"
    normalized = normalize_text(raw)
    assert (
        normalized
        == "Proficient in C++, C#, .NET 8, Node.js, JavaScript, and SQL.\nWebsite: https://nextround.in\nEmail: contact@nextround.in"
    )


def test_normalize_text_preserves_paragraph_line_structure() -> None:
    raw = "John Doe\nSoftware Engineer\n\nSkills:\nPython, FastAPI\n\nExperience:\nBackend Intern"
    normalized = normalize_text(raw)
    assert (
        normalized
        == "John Doe\nSoftware Engineer\n\nSkills:\nPython, FastAPI\n\nExperience:\nBackend Intern"
    )


def test_normalize_empty_text() -> None:
    assert normalize_text("") == ""
    assert normalize_text("   \n\n \t  \r\n  ") == ""


# ==========================================
# 2. PDF EXTRACTOR TESTS
# ==========================================


def test_pdf_extract_single_page(tmp_path: Path) -> None:
    pdf_path = tmp_path / "single_page.pdf"
    pdf_path.write_bytes(create_pdf_bytes(["Alice Smith - Backend Developer"]))

    extractor = PDFExtractor()
    assert extractor.supports("application/pdf")
    assert extractor.supports("application/PDF")
    assert extractor.supports("", ".pdf")

    result = extractor.extract(pdf_path)
    assert isinstance(result, ExtractionResult)
    assert result.extractor_type == "pdf"
    assert result.status == ExtractionStatus.COMPLETE
    assert result.extraction_issues == []
    assert result.page_count == 1
    assert "Alice Smith - Backend Developer" in result.text
    assert result.character_count == len(result.text)


def test_pdf_extract_multi_page_in_order(tmp_path: Path) -> None:
    pdf_path = tmp_path / "multi_page.pdf"
    pdf_path.write_bytes(
        create_pdf_bytes(
            [
                "Page 1: Education Summary",
                "Page 2: Projects and Work Experience",
                "Page 3: Certifications and Achievements",
            ]
        )
    )

    extractor = PDFExtractor()
    result = extractor.extract(pdf_path)
    assert result.status == ExtractionStatus.COMPLETE
    assert result.extraction_issues == []
    assert result.page_count == 3

    lines = result.text.split("\n\n")
    assert len(lines) == 3
    assert lines[0] == "Page 1: Education Summary"
    assert lines[1] == "Page 2: Projects and Work Experience"
    assert lines[2] == "Page 3: Certifications and Achievements"


def test_pdf_extract_complete_with_intentionally_textless_page(tmp_path: Path) -> None:
    # A multi-page PDF where one page is textless (e.g. image/blank) is COMPLETE if all pages processed without error
    pdf_path = tmp_path / "partially_textless.pdf"
    pdf_path.write_bytes(
        create_pdf_bytes(["Page 1: Text Content", "", "Page 3: More Text"])
    )

    extractor = PDFExtractor()
    result = extractor.extract(pdf_path)
    assert result.status == ExtractionStatus.COMPLETE
    assert result.extraction_issues == []
    assert result.page_count == 3
    assert "Page 1: Text Content\n\nPage 3: More Text" == result.text


def test_pdf_extract_empty_scanned_pdf(tmp_path: Path) -> None:
    # A valid image-only/scanned PDF without embedded text stream produces EMPTY status and no issues
    pdf_path = tmp_path / "scanned.pdf"
    pdf_path.write_bytes(create_pdf_bytes([""]))

    extractor = PDFExtractor()
    result = extractor.extract(pdf_path)
    assert result.status == ExtractionStatus.EMPTY
    assert result.extraction_issues == []
    assert result.page_count == 1
    assert result.text == ""
    assert result.character_count == 0


def test_pdf_extract_partial_when_single_page_fails(tmp_path: Path) -> None:
    pdf_path = tmp_path / "multi_page_partial.pdf"
    pdf_path.write_bytes(
        create_pdf_bytes(["Page 1: Extractable info", "Page 2: Second page info"])
    )

    extractor = PDFExtractor()
    with patch("pdfplumber.open") as mock_open:
        mock_pdf = MagicMock()
        mock_page1 = MagicMock()
        mock_page1.extract_text.return_value = "Page 1: Extractable info"
        mock_page2 = MagicMock()
        mock_page2.extract_text.side_effect = RuntimeError(
            "Corrupt font table on page 2"
        )

        mock_pdf.pages = [mock_page1, mock_page2]
        mock_open.return_value.__enter__.return_value = mock_pdf

        result = extractor.extract(pdf_path)
        assert result.status == ExtractionStatus.PARTIAL
        assert len(result.extraction_issues) == 1
        issue = result.extraction_issues[0]
        assert isinstance(issue, ExtractionIssue)
        assert issue.location_type == "page"
        assert issue.location == 2
        assert "Corrupt font table on page 2" in issue.reason
        assert result.page_count == 2
        assert "Page 1: Extractable info" in result.text
        assert result.character_count == len(result.text)


def test_pdf_extract_nonexistent_file(tmp_path: Path) -> None:
    extractor = PDFExtractor()
    with pytest.raises(DocumentReadError) as exc_info:
        extractor.extract(tmp_path / "nonexistent.pdf")
    assert "not found" in str(exc_info.value).lower()


def test_pdf_extract_corrupt_file(tmp_path: Path) -> None:
    pdf_path = tmp_path / "corrupt.pdf"
    pdf_path.write_bytes(b"%PDF-1.4\nCorrupted binary payload that is unparseable")

    extractor = PDFExtractor()
    with pytest.raises(DocumentContentError):
        extractor.extract(pdf_path)


# ==========================================
# 3. DOCX EXTRACTOR TESTS
# ==========================================


def test_docx_extract_paragraphs(tmp_path: Path) -> None:
    docx_path = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_paragraph("Bob Jones")
    doc.add_paragraph("Frontend Engineer")
    doc.add_paragraph("Skills: React, Next.js, TypeScript")
    doc.save(docx_path)

    extractor = DOCXExtractor()
    assert extractor.supports(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert extractor.supports("", ".docx")

    result = extractor.extract(docx_path)
    assert result.extractor_type == "docx"
    assert result.status == ExtractionStatus.COMPLETE
    assert result.extraction_issues == []
    assert result.paragraph_count == 3
    assert (
        result.text
        == "Bob Jones\n\nFrontend Engineer\n\nSkills: React, Next.js, TypeScript"
    )
    assert result.character_count == len(result.text)


def test_docx_extract_empty_document(tmp_path: Path) -> None:
    docx_path = tmp_path / "empty.docx"
    doc = docx.Document()
    doc.save(docx_path)

    extractor = DOCXExtractor()
    result = extractor.extract(docx_path)
    assert result.status == ExtractionStatus.EMPTY
    assert result.extraction_issues == []
    assert result.text == ""
    assert result.character_count == 0


def test_docx_extract_complete_with_intentionally_empty_paragraph(
    tmp_path: Path,
) -> None:
    docx_path = tmp_path / "with_empty_p.docx"
    doc = docx.Document()
    doc.add_paragraph("Header Title")
    doc.add_paragraph("")  # empty paragraph
    doc.add_paragraph("Body Content")
    doc.save(docx_path)

    extractor = DOCXExtractor()
    result = extractor.extract(docx_path)
    assert result.status == ExtractionStatus.COMPLETE
    assert result.extraction_issues == []
    assert result.text == "Header Title\n\nBody Content"


def test_docx_extract_tables_and_paragraphs_reading_order(tmp_path: Path) -> None:
    docx_path = tmp_path / "with_table.docx"
    doc = docx.Document()
    doc.add_paragraph("Resume Title: Engineering Profile")

    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Institution"
    table.cell(0, 1).text = "Degree"
    table.cell(1, 0).text = "IIT Bombay"
    table.cell(1, 1).text = "B.Tech Computer Science"

    doc.add_paragraph("Projects: Placement Readiness Portal")
    doc.save(docx_path)

    extractor = DOCXExtractor()
    result = extractor.extract(docx_path)

    expected = (
        "Resume Title: Engineering Profile\n\n"
        "Institution Degree\n"
        "IIT Bombay B.Tech Computer Science\n\n"
        "Projects: Placement Readiness Portal"
    )
    assert result.status == ExtractionStatus.COMPLETE
    assert result.extraction_issues == []
    assert result.text == expected


def test_docx_extract_partial_when_element_fails(tmp_path: Path) -> None:
    docx_path = tmp_path / "partial.docx"
    doc = docx.Document()
    doc.add_paragraph("Valid Intro Paragraph")
    doc.save(docx_path)

    extractor = DOCXExtractor()
    with patch("docx.Document") as mock_doc_ctor:
        mock_doc = MagicMock()
        mock_child1 = MagicMock()
        mock_child1.tag = (
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
        )
        mock_child2 = MagicMock()
        mock_child2.tag = (
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl"
        )

        mock_doc.element.body = [mock_child1, mock_child2]

        # First element succeeds as Paragraph, second fails as Table
        with (
            patch("app.services.extraction.docx.Paragraph") as mock_p,
            patch("app.services.extraction.docx.Table") as mock_t,
        ):
            mock_p_instance = MagicMock()
            mock_p_instance.text = "Extracted paragraph content"
            mock_p.return_value = mock_p_instance

            mock_t.side_effect = RuntimeError("Corrupt XML in table element")
            mock_doc_ctor.return_value = mock_doc

            result = extractor.extract(docx_path)
            assert result.status == ExtractionStatus.PARTIAL
            assert len(result.extraction_issues) == 1
            issue = result.extraction_issues[0]
            assert isinstance(issue, ExtractionIssue)
            assert issue.location_type == "table"
            assert issue.location == 1
            assert "Corrupt XML in table element" in issue.reason
            assert "Extracted paragraph content" in result.text
            assert result.character_count == len(result.text)


def test_docx_extract_nonexistent_file(tmp_path: Path) -> None:
    extractor = DOCXExtractor()
    with pytest.raises(DocumentReadError) as exc_info:
        extractor.extract(tmp_path / "nonexistent.docx")
    assert "not found" in str(exc_info.value).lower()


def test_docx_extract_corrupt_file(tmp_path: Path) -> None:
    docx_path = tmp_path / "corrupt.docx"
    docx_path.write_bytes(b"PK\x03\x04\nCorrupt DOCX package that is not valid zip")

    extractor = DOCXExtractor()
    with pytest.raises(DocumentContentError):
        extractor.extract(docx_path)


# ==========================================
# 4. REGISTRY TESTS
# ==========================================


def test_registry_selection() -> None:
    registry = ExtractorRegistry()

    # PDF selection
    pdf_ext = registry.get_extractor("application/pdf")
    assert isinstance(pdf_ext, PDFExtractor)
    assert pdf_ext.extractor_type == "pdf"

    # DOCX selection
    docx_ext = registry.get_extractor(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert isinstance(docx_ext, DOCXExtractor)
    assert docx_ext.extractor_type == "docx"

    # Extension fallback
    pdf_by_ext = registry.get_extractor("", ".pdf")
    assert isinstance(pdf_by_ext, PDFExtractor)


def test_registry_unsupported_format_raises() -> None:
    registry = ExtractorRegistry()

    with pytest.raises(UnsupportedDocumentError):
        registry.get_extractor("text/plain", ".txt")

    with pytest.raises(UnsupportedDocumentError):
        registry.get_extractor("application/msword", ".doc")

    with pytest.raises(UnsupportedDocumentError):
        registry.get_extractor("image/png", ".png")


# ==========================================
# 5. STORAGE & SERVICE INTEGRATION TESTS
# ==========================================


def test_document_extraction_service_with_resume(tmp_path: Path) -> None:
    storage = LocalStorageService(root_dir=tmp_path)
    service = DocumentExtractionService(storage_service=storage)

    user_id = uuid.uuid4()
    resume_id = uuid.uuid4()
    storage_key = f"users/{user_id}/resumes/{resume_id}.pdf"

    # Save PDF into local storage
    target_path = storage.resolve_path(storage_key)
    target_path.parent.mkdir(parents=True, exist_ok=True)
    target_path.write_bytes(
        create_pdf_bytes(["Test Candidate Name\nFull Stack Developer"])
    )

    resume = Resume(
        id=resume_id,
        user_id=user_id,
        original_filename="../../malicious/attempt.pdf",  # Must never be used as storage path
        storage_key=storage_key,
        mime_type="application/pdf",
        file_size=1024,
        version=1,
        status=ResumeStatus.UPLOADED,
    )

    result = service.extract_from_resume(resume)
    assert result.extractor_type == "pdf"
    assert result.status == ExtractionStatus.COMPLETE
    assert result.extraction_issues == []
    assert "Test Candidate Name\nFull Stack Developer" in result.text
    assert result.character_count == len(result.text)


def test_document_extraction_service_nonexistent_storage_file(tmp_path: Path) -> None:
    storage = LocalStorageService(root_dir=tmp_path)
    service = DocumentExtractionService(storage_service=storage)

    resume = Resume(
        id=uuid.uuid4(),
        user_id=uuid.uuid4(),
        original_filename="missing.pdf",
        storage_key="users/missing/resumes/missing.pdf",
        mime_type="application/pdf",
        file_size=1000,
        version=1,
    )

    with pytest.raises(DocumentReadError):
        service.extract_from_resume(resume)


def test_document_extraction_service_path_traversal_blocked(tmp_path: Path) -> None:
    storage = LocalStorageService(root_dir=tmp_path)
    service = DocumentExtractionService(storage_service=storage)

    resume = Resume(
        id=uuid.uuid4(),
        user_id=uuid.uuid4(),
        original_filename="escape.pdf",
        storage_key="../../../etc/passwd",
        mime_type="application/pdf",
        file_size=1000,
        version=1,
    )

    with pytest.raises(StoragePathTraversalError):
        service.extract_from_resume(resume)

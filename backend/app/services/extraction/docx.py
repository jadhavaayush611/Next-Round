from pathlib import Path

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.services.extraction.base import BaseDocumentExtractor
from app.services.extraction.exceptions import DocumentContentError, DocumentReadError
from app.services.extraction.models import (
    ExtractedDocument,
    ExtractedElement,
    ExtractionIssue,
    ExtractionResult,
    ExtractionStatus,
)
from app.services.extraction.normalizer import normalize_text


class DOCXExtractor(BaseDocumentExtractor):
    """Local DOCX text extractor utilizing python-docx."""

    SUPPORTED_MIME_TYPES = {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    }
    SUPPORTED_EXTENSIONS = {".docx"}

    @property
    def extractor_type(self) -> str:
        return "docx"

    def supports(self, mime_type: str, extension: str | None = None) -> bool:
        norm_mime = (mime_type or "").lower().split(";")[0].strip()
        if norm_mime in self.SUPPORTED_MIME_TYPES:
            return True
        if extension:
            norm_ext = (
                extension.lower()
                if extension.startswith(".")
                else f".{extension.lower()}"
            )
            if norm_ext in self.SUPPORTED_EXTENSIONS:
                return True
        return False

    def extract(self, file_path: Path | str) -> ExtractionResult:
        path = Path(file_path)
        if not path.exists() or not path.is_file():
            raise DocumentReadError(f"DOCX document not found at: {path}")

        try:
            doc = docx.Document(str(path))
            content_blocks: list[str] = []
            issues: list[ExtractionIssue] = []
            paragraph_count = 0
            table_count = 0
            element_index = 0

            # Iterate child elements in document body to preserve exact reading order
            for child in doc.element.body:
                element_index += 1
                try:
                    if child.tag.endswith("p"):
                        paragraph_count += 1
                        p = Paragraph(child, doc)
                        if p.text.strip():
                            content_blocks.append(p.text)
                    elif child.tag.endswith("tbl"):
                        table_count += 1
                        t = Table(child, doc)
                        table_rows: list[str] = []
                        for row in t.rows:
                            seen_cells: set[object] = set()
                            cell_texts: list[str] = []
                            for cell in row.cells:
                                if cell._tc not in seen_cells:
                                    seen_cells.add(cell._tc)
                                    cell_txt = cell.text.strip()
                                    if cell_txt:
                                        cell_texts.append(cell_txt)
                            if cell_texts:
                                table_rows.append("\t".join(cell_texts))
                        if table_rows:
                            content_blocks.append("\n".join(table_rows))
                except Exception as exc:
                    tag_name = (
                        child.tag.split("}")[-1] if "}" in child.tag else child.tag
                    )
                    if tag_name == "p":
                        loc_type = "paragraph"
                        loc = paragraph_count if paragraph_count > 0 else element_index
                    elif tag_name == "tbl":
                        loc_type = "table"
                        loc = table_count if table_count > 0 else element_index
                    else:
                        loc_type = "element"
                        loc = element_index

                    issues.append(
                        ExtractionIssue(
                            location_type=loc_type,
                            location=loc,
                            reason=f"Failed to extract text from {loc_type} at index {loc}: {exc}",
                        )
                    )

            raw_text = "\n\n".join(content_blocks)
            normalized = normalize_text(raw_text)

            if issues:
                if normalized:
                    status = ExtractionStatus.PARTIAL
                else:
                    raise DocumentContentError(
                        "All DOCX body elements failed to extract content."
                    )
            else:
                if normalized:
                    status = ExtractionStatus.COMPLETE
                else:
                    status = ExtractionStatus.EMPTY

            return ExtractionResult(
                text=normalized,
                character_count=len(normalized),
                extractor_type=self.extractor_type,
                status=status,
                paragraph_count=paragraph_count,
                extraction_issues=issues,
            )
        except (FileNotFoundError, PermissionError) as exc:
            raise DocumentReadError(f"Failed to access DOCX document: {exc}") from None
        except DocumentContentError:
            raise
        except Exception as exc:
            raise DocumentContentError(
                f"Failed to extract text from DOCX document: {exc}"
            ) from None

    def extract_document(self, file_path: Path | str) -> ExtractedDocument:
        path = Path(file_path)
        if not path.exists() or not path.is_file():
            raise DocumentReadError(f"DOCX document not found at: {path}")

        try:
            doc = docx.Document(str(path))
            elements: list[ExtractedElement] = []
            content_blocks: list[str] = []
            issues: list[ExtractionIssue] = []
            paragraph_count = 0
            table_count = 0
            element_index = 0
            elem_order = 0

            for child in doc.element.body:
                element_index += 1
                try:
                    if child.tag.endswith("p"):
                        paragraph_count += 1
                        p = Paragraph(child, doc)
                        txt = p.text.strip()
                        if txt:
                            content_blocks.append(p.text)
                            elem_order += 1
                            elements.append(
                                ExtractedElement(
                                    element_id=elem_order,
                                    page_number=None,
                                    order=elem_order,
                                    element_type="paragraph",
                                    text=normalize_text(txt),
                                )
                            )
                    elif child.tag.endswith("tbl"):
                        table_count += 1
                        t = Table(child, doc)
                        table_rows: list[str] = []
                        for row in t.rows:
                            seen_cells: set[object] = set()
                            cell_texts: list[str] = []
                            for cell in row.cells:
                                if cell._tc not in seen_cells:
                                    seen_cells.add(cell._tc)
                                    cell_txt = cell.text.strip()
                                    if cell_txt:
                                        cell_texts.append(cell_txt)
                            if cell_texts:
                                table_rows.append("\t".join(cell_texts))
                        if table_rows:
                            table_full = "\n".join(table_rows)
                            content_blocks.append(table_full)
                            elem_order += 1
                            elements.append(
                                ExtractedElement(
                                    element_id=elem_order,
                                    page_number=None,
                                    order=elem_order,
                                    element_type="table",
                                    text=normalize_text(table_full),
                                )
                            )
                except Exception as exc:
                    tag_name = (
                        child.tag.split("}")[-1] if "}" in child.tag else child.tag
                    )
                    if tag_name == "p":
                        loc_type = "paragraph"
                        loc = paragraph_count if paragraph_count > 0 else element_index
                    elif tag_name == "tbl":
                        loc_type = "table"
                        loc = table_count if table_count > 0 else element_index
                    else:
                        loc_type = "element"
                        loc = element_index

                    issues.append(
                        ExtractionIssue(
                            location_type=loc_type,
                            location=loc,
                            reason=f"Failed to extract text from {loc_type} at index {loc}: {exc}",
                        )
                    )

            raw_text = "\n\n".join(content_blocks)
            normalized = normalize_text(raw_text)

            if issues and not normalized:
                raise DocumentContentError(
                    "All DOCX body elements failed to extract content."
                )

            metadata = {
                "extractor_type": self.extractor_type,
                "paragraph_count": paragraph_count,
                "table_count": table_count,
                "character_count": len(normalized),
            }

            return ExtractedDocument(
                elements=elements,
                full_text=normalized,
                extraction_issues=issues,
                metadata=metadata,
            )
        except (FileNotFoundError, PermissionError) as exc:
            raise DocumentReadError(f"Failed to access DOCX document: {exc}") from None
        except DocumentContentError:
            raise
        except Exception as exc:
            raise DocumentContentError(
                f"Failed to extract text from DOCX document: {exc}"
            ) from None
